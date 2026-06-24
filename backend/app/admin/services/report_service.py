"""Reporting business rules: analytics aggregates and data exports.

Builds tabular datasets from the existing tables and serialises them to CSV or
XLSX in-memory (no files written to disk, no schema changes). The summary
endpoint returns headline numbers for the Reports dashboards.
"""

from __future__ import annotations

import csv
import io

from app.admin.cache import ttl_cached
from app.admin.repositories import (
    CandidateRepository,
    HrRepository,
    JobPoolRepository,
)
from app.admin.services.application_service import ApplicationService
from app.admin.services.company_service import CompanyService

# Column order for each exportable dataset. Keys must exist on the row dicts.
_DATASETS = {
    "candidates": ["id", "full_name", "email", "phone", "city", "current_job_title",
                   "current_company", "years_of_experience", "open_to_work", "created_at"],
    "recruiters": ["id", "full_name", "email", "phone", "company_name", "company_industry",
                   "company_size", "created_at"],
    "companies": ["company_name", "recruiters", "jobs", "company_industry", "company_size",
                  "company_website", "company_email"],
    "jobs": ["id", "title", "company_name", "location", "contract_type", "seniority_level",
             "status", "archived", "created_at"],
    "applications": ["session_id", "candidate_name", "pool_title", "status", "answered",
                     "total_questions", "updated_at"],
}
# "users" is the union of candidates + recruiters.
_DATASETS["users"] = ["id", "type", "full_name", "email", "phone", "company_name", "created_at"]


class ReportService:
    def __init__(self) -> None:
        self.candidates = CandidateRepository()
        self.recruiters = HrRepository()
        self.pools = JobPoolRepository()

    # --- analytics summary ---------------------------------------------------

    @ttl_cached(20.0)
    def summary(self) -> dict:
        candidates = self.candidates.list_all()
        recruiters = self.recruiters.list_all()
        pools = self.pools.list_all()
        applications = ApplicationService().list_applications()
        companies = CompanyService().list_companies()

        completed = sum(1 for a in applications if a["status"] == "completed")
        open_to_work = sum(1 for c in candidates if c.get("open_to_work"))
        completion_rate = round((completed / len(applications) * 100), 1) if applications else 0.0

        return {
            "totals": {
                "candidates": len(candidates),
                "recruiters": len(recruiters),
                "companies": len(companies),
                "jobs": len(pools),
                "activeJobs": sum(1 for p in pools if p.get("status")),
                "archivedJobs": sum(1 for p in pools if p.get("archived")),
                "applications": len(applications),
                "completedApplications": completed,
            },
            "rates": {
                "interviewCompletionRate": completion_rate,
                "openToWorkCandidates": open_to_work,
            },
            "topCompanies": sorted(companies, key=lambda c: c["jobs"], reverse=True)[:5],
        }

    # --- export --------------------------------------------------------------

    def build_rows(self, dataset: str) -> tuple[list[str], list[dict]]:
        """Return (columns, rows) for a dataset name."""
        columns = _DATASETS[dataset]
        if dataset == "candidates":
            rows = self.candidates.list_all()
        elif dataset == "recruiters":
            rows = self.recruiters.list_all()
        elif dataset == "jobs":
            rows = JobServiceRows(self.pools)
        elif dataset == "companies":
            rows = CompanyService().list_companies()
        elif dataset == "applications":
            rows = ApplicationService().list_applications()
        elif dataset == "users":
            rows = self._user_rows()
        else:  # pragma: no cover -- guarded by validator upstream
            rows = []
        return columns, list(rows)

    def _user_rows(self) -> list[dict]:
        rows: list[dict] = []
        for c in self.candidates.list_all():
            rows.append({**c, "type": "candidate"})
        for r in self.recruiters.list_all():
            rows.append({**r, "type": "recruiter"})
        return rows

    def to_csv(self, dataset: str) -> bytes:
        columns, rows = self.build_rows(dataset)
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: _cell(row.get(k)) for k in columns})
        return buf.getvalue().encode("utf-8-sig")  # BOM => Excel opens UTF-8 cleanly

    def to_xlsx(self, dataset: str) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        columns, rows = self.build_rows(dataset)
        wb = Workbook()
        ws = wb.active
        ws.title = dataset[:31] or "report"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F46E5")  # indigo-600
        ws.append([c.replace("_", " ").title() for c in columns])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for row in rows:
            ws.append([_cell(row.get(k)) for k in columns])

        for i, col in enumerate(columns, start=1):
            width = max(len(col) + 2, 14)
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 40)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()


def JobServiceRows(pool_repo) -> list[dict]:  # noqa: N802 -- tiny adapter, keeps export uniform
    """Flatten job pools (with company) for the jobs export."""
    from app.admin.services.job_service import _flatten
    return [_flatten(p) for p in pool_repo.list_with_recruiter()]


def _cell(value):
    """Render a value as an export-friendly scalar."""
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value
